import pandas as pd
import os
import re
from openpyxl import load_workbook

def add_prefix_to_excel(file_path, prefix):
    """
    在 Excel 文件的文件名列添加前缀并更新新链接列
    :param file_path: Excel 文件路径
    :param prefix: 前缀
    :return: 更新后的 Excel 文件路径
    """
    try:
        df = pd.read_excel(file_path)
    except FileNotFoundError:
        print(f"未找到文件: {file_path}，请检查文件路径是否正确。")
        return None

    # 在文件名列添加前缀并更新新链接列
    df['新链接'] = prefix + df['文件名']

    # 保存修改后的 DataFrame 到新的 Excel 文件
    new_file_path = file_path.rsplit('.', 1)[0] + '_new.xlsx'
    df.to_excel(new_file_path, index=False)
    print(f"处理完成，新文件已保存到: {new_file_path}")
    return new_file_path

def load_desensitization_mapping(excel_file):
    """
    从 Excel 文件中加载原链接和新链接的映射关系
    :param excel_file: Excel 文件路径
    :return: 原链接到新链接的映射字典
    """
    wb = load_workbook(excel_file)
    ws = wb.active
    mapping = {}
    # 获取表头行
    header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    original_link_col = None
    new_link_col = None
    # 查找原链接和新链接所在列索引
    for col_index, header in enumerate(header_row):
        if header == '原链接':
            original_link_col = col_index
        elif header == '新链接':
            new_link_col = col_index

    if original_link_col is None or new_link_col is None:
        print("Excel 文件中未找到 '原链接' 或 '新链接' 列。")
        return {}

    # 从第二行开始遍历 Excel 表格，第一行通常是表头
    for row in ws.iter_rows(min_row=2, values_only=True):
        original_link = row[original_link_col]
        new_link = row[new_link_col]
        if original_link and new_link:
            mapping[original_link] = new_link
    return mapping

def desensitize_md_files(folder_path, mapping):
    """
    对指定文件夹下的所有 .md 文件进行链接替换处理
    :param folder_path: 包含 .md 文件的文件夹路径
    :param mapping: 原链接到新链接的映射字典
    """
    # 用于记录每个原链接是否被替换，初始都设为 False
    replaced_status = {link: False for link in mapping}

    for root, _, files in os.walk(folder_path):
        for file in files:
            if file.endswith('.md'):
                file_path = os.path.join(root, file)
                with open(file_path, 'r', encoding='utf-8') as f:
                    content = f.read()

                for original_link in mapping:
                    pattern = re.escape(original_link)
                    new_content, count = re.subn(pattern, mapping[original_link], content)
                    if count > 0:
                        # 如果替换成功，将该原链接的替换状态设为 True
                        replaced_status[original_link] = True
                    content = new_content

                with open(file_path, 'w', encoding='utf-8') as f:
                    f.write(content)

    # 找出未被替换的原链接
    not_replaced_links = []
    for original_link, replaced in replaced_status.items():
        if not replaced:
            not_replaced_links.append((original_link, mapping[original_link]))

    if not_replaced_links:
        print("以下部分未成功替换：")
        for original_link, new_link in not_replaced_links:
            print(f"文件链接清单 - {original_link} -> {new_link}")
    else:
        print("所有链接都已成功替换。")

if __name__ == "__main__":
    # 获取用户输入
    file_path = input("请输入链接清单 Excel 文件的本地存储地址: ")
    prefix = input("请输入新链接前缀: ")
    folder_path = input("请输入 Markdown 文档所在 docs 文件夹地址: ")

    # 在 Excel 文件的文件名列添加前缀并更新新链接列
    new_excel_file = add_prefix_to_excel(file_path, prefix)
    if new_excel_file:
        # 加载原链接和新链接的映射关系
        mapping = load_desensitization_mapping(new_excel_file)
        # 对 .md 文件进行链接替换处理
        desensitize_md_files(folder_path, mapping)
        print("链接替换处理完成。")